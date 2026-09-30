import os
import json
import time
import random
from datetime import datetime

import openpyxl
from playwright.sync_api import sync_playwright


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Excel input file
CONTACTS_FILE = os.path.join(
    BASE_DIR,
    "contacts.xlsx"
)

# Persistent Chrome / WhatsApp profile
SESSION_DIR = os.path.join(
    BASE_DIR,
    "whatsapp_session"
)

# Screenshots
SCREENSHOT_DIR = os.path.join(
    BASE_DIR,
    "screenshots"
)

# Reports
REPORT_DIR = os.path.join(
    BASE_DIR,
    "reports"
)

# Your existing Google Chrome
CHROME_PATH = (
    r"C:\Program Files\Google\Chrome\Application\chrome.exe"
)

# Browser should be visible
HEADLESS = False

# Assignment requirement:
# random delays between actions
MIN_DELAY = 2
MAX_DELAY = 5


# ============================================================
# CREATE DIRECTORIES
# ============================================================

os.makedirs(
    SESSION_DIR,
    exist_ok=True
)

os.makedirs(
    SCREENSHOT_DIR,
    exist_ok=True
)

os.makedirs(
    REPORT_DIR,
    exist_ok=True
)


# ============================================================
# RANDOM HUMAN-LIKE DELAY
# ============================================================

def human_delay(
    minimum=MIN_DELAY,
    maximum=MAX_DELAY
):
    delay = random.uniform(
        minimum,
        maximum
    )

    print(
        f"Waiting {delay:.2f} seconds..."
    )

    time.sleep(delay)


# ============================================================
# CLEAN PHONE NUMBER
# ============================================================

def clean_phone_number(phone):

    phone = str(phone).strip()

    phone = (
        phone
        .replace(" ", "")
        .replace("-", "")
        .replace("(", "")
        .replace(")", "")
    )

    return phone


# ============================================================
# CREATE CONTACTS.XLSX
# ============================================================

def create_contacts_file():

    if os.path.exists(
        CONTACTS_FILE
    ):

        print(
            f"Found contacts.xlsx:"
        )

        print(
            CONTACTS_FILE
        )

        return

    print()
    print("=" * 60)
    print("contacts.xlsx not found.")
    print("Creating sample contacts.xlsx...")
    print("=" * 60)
    print()

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "Contacts"

    # Required columns
    sheet.append([
        "Name",
        "Phone",
        "Message"
    ])

    # --------------------------------------------------------
    # SAMPLE CONTACT
    #
    # IMPORTANT:
    # Replace this number with a real test number
    # before allowing the automation to send.
    # --------------------------------------------------------

    sheet.append([
        "Test User",
        "+919999999999",
        "Hello {name}, this is a Playwright automation test."
    ])

    workbook.save(
        CONTACTS_FILE
    )

    print(
        f"Created: {CONTACTS_FILE}"
    )

    print()
    print(
        "IMPORTANT:"
    )

    print(
        "Open contacts.xlsx and replace the"
    )

    print(
        "sample phone number with your test number."
    )

    print()


# ============================================================
# LOAD CONTACTS
# ============================================================

def load_contacts():

    if not os.path.exists(
        CONTACTS_FILE
    ):

        raise FileNotFoundError(
            f"contacts.xlsx not found at: "
            f"{CONTACTS_FILE}"
        )

    workbook = openpyxl.load_workbook(
        CONTACTS_FILE,
        data_only=True
    )

    sheet = workbook.active

    headers = [
        cell.value
        for cell in sheet[1]
    ]

    required_columns = [
        "Name",
        "Phone",
        "Message"
    ]

    for column in required_columns:

        if column not in headers:

            raise ValueError(
                f"Missing required column: {column}"
            )

    name_index = headers.index(
        "Name"
    )

    phone_index = headers.index(
        "Phone"
    )

    message_index = headers.index(
        "Message"
    )

    contacts = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not row:
            continue

        name = row[name_index]

        phone = row[phone_index]

        message = row[message_index]

        # Skip incomplete rows
        if not name or not phone:
            continue

        name = str(
            name
        ).strip()

        phone = clean_phone_number(
            phone
        )

        if message:

            message = str(
                message
            ).strip()

        else:

            message = "Hello {name}"

        # Personalize message
        message = message.replace(
            "{name}",
            name
        )

        contacts.append({

            "name":
                name,

            "phone":
                phone,

            "message":
                message
        })

    return contacts


# ============================================================
# WAIT FOR WHATSAPP LOGIN
# ============================================================

def wait_for_whatsapp_login(page):

    print()
    print("=" * 60)
    print("WHATSAPP WEB LOGIN")
    print("=" * 60)

    print()
    print(
        "If the QR code is displayed:"
    )

    print(
        "WhatsApp -> Linked Devices -> Link a Device"
    )

    print()
    print(
        "Waiting for WhatsApp Web..."
    )

    try:

        page.locator(
            "#pane-side"
        ).wait_for(
            state="visible",
            timeout=120000
        )

        print()
        print(
            "WhatsApp Web login detected!"
        )

        return True

    except Exception as error:

        print()
        print(
            "WhatsApp login was not detected."
        )

        print(
            f"Error: {error}"
        )

        return False


# ============================================================
# OPEN CONTACT
# ============================================================

def open_contact(
    page,
    phone
):

    phone_number = (
        phone
        .replace("+", "")
        .replace(" ", "")
        .replace("-", "")
        .replace("(", "")
        .replace(")", "")
    )

    url = (
        "https://web.whatsapp.com/send"
        f"?phone={phone_number}"
    )

    print()
    print(
        f"Opening WhatsApp chat for: {phone}"
    )

    page.goto(
        url,
        wait_until="domcontentloaded"
    )

    page.wait_for_timeout(
        3000
    )

    human_delay(
        2,
        4
    )


# ============================================================
# FIND MESSAGE BOX
# ============================================================

def get_message_box(page):

    selectors = [

        'div[contenteditable="true"][role="textbox"]',

        'div[contenteditable="true"]'

    ]

    for selector in selectors:

        try:

            elements = page.locator(
                selector
            )

            count = elements.count()

            if count == 0:
                continue

            # Search from the last element backwards
            for index in range(
                count - 1,
                -1,
                -1
            ):

                element = elements.nth(
                    index
                )

                try:

                    if element.is_visible(
                        timeout=2000
                    ):

                        return element

                except Exception:

                    continue

        except Exception:

            continue

    return None


# ============================================================
# SEND MESSAGE
# ============================================================

def send_message(
    page,
    message
):

    print()
    print(
        "Preparing message..."
    )

    print(
        f"Message: {message}"
    )

    message_box = get_message_box(
        page
    )

    if message_box is None:

        raise RuntimeError(
            "WhatsApp message box was not found."
        )

    human_delay(
        2,
        4
    )

    message_box.click()

    try:

        message_box.fill(
            ""
        )

    except Exception:

        pass

    message_box.fill(
        message
    )

    print(
        "Message typed."
    )

    human_delay(
        1,
        3
    )

    message_box.press(
        "Enter"
    )

    print(
        "Message sent."
    )

    human_delay(
        2,
        5
    )


# ============================================================
# SCREENSHOT
# ============================================================

def take_screenshot(
    page,
    name
):

    safe_name = "".join(

        character
        if character.isalnum()
        else "_"

        for character in name
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        f"{safe_name}_{timestamp}.png"
    )

    filepath = os.path.join(
        SCREENSHOT_DIR,
        filename
    )

    page.screenshot(
        path=filepath,
        full_page=False
    )

    print()
    print(
        f"Screenshot saved:"
    )

    print(
        filepath
    )

    return filepath


# ============================================================
# EXTRACT LAST 3 MESSAGES
# ============================================================

def extract_last_three_messages(page):

    print()
    print(
        "Extracting last 3 messages..."
    )

    messages = []

    try:

        page.wait_for_timeout(
            2000
        )

        message_elements = page.locator(
            '[data-pre-plain-text]'
        )

        count = message_elements.count()

        print(
            f"Message elements found: {count}"
        )

        if count == 0:

            print(
                "No message elements found."
            )

            return []

        start_index = max(
            0,
            count - 3
        )

        for index in range(
            start_index,
            count
        ):

            element = (
                message_elements
                .nth(index)
            )

            try:

                text = element.inner_text(
                    timeout=3000
                )

                text = text.strip()

                if text:

                    messages.append(
                        text
                    )

            except Exception:

                continue

    except Exception as error:

        print(
            f"Message extraction error: {error}"
        )

    messages = messages[-3:]

    print(
        f"Extracted {len(messages)} message(s)."
    )

    return messages


# ============================================================
# SAVE JSON REPORT
# ============================================================

def save_json_report(
    results
):

    today = datetime.now().strftime(
        "%Y-%m-%d"
    )

    filename = (
        f"whatsapp_report_{today}.json"
    )

    filepath = os.path.join(
        REPORT_DIR,
        filename
    )

    successful = sum(

        1

        for result in results

        if result["status"] == "SENT"
    )

    failed = sum(

        1

        for result in results

        if result["status"] == "FAILED"
    )

    report = {

        "generated_at":
            datetime.now().isoformat(),

        "total_contacts":
            len(results),

        "successful":
            successful,

        "failed":
            failed,

        "results":
            results
    }

    with open(
        filepath,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            report,
            file,
            indent=4,
            ensure_ascii=False
        )

    print()
    print(
        f"JSON report saved:"
    )

    print(
        filepath
    )


# ============================================================
# SAVE EXCEL REPORT
# ============================================================

def save_excel_report(
    results
):

    today = datetime.now().strftime(
        "%Y-%m-%d"
    )

    filename = (
        f"whatsapp_report_{today}.xlsx"
    )

    filepath = os.path.join(
        REPORT_DIR,
        filename
    )

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "WhatsApp Report"

    headers = [

        "Name",

        "Phone",

        "Message",

        "Status",

        "Error",

        "Last 3 Messages",

        "Screenshot",

        "Processed At"
    ]

    sheet.append(
        headers
    )

    for result in results:

        messages = result.get(
            "last_3_messages",
            []
        )

        messages_text = "\n".join(
            messages
        )

        sheet.append([

            result.get(
                "name",
                ""
            ),

            result.get(
                "phone",
                ""
            ),

            result.get(
                "message",
                ""
            ),

            result.get(
                "status",
                ""
            ),

            result.get(
                "error",
                ""
            ),

            messages_text,

            result.get(
                "screenshot",
                ""
            ),

            result.get(
                "processed_at",
                ""
            )
        ])

    # Column widths
    widths = {

        "A": 20,

        "B": 20,

        "C": 50,

        "D": 15,

        "E": 50,

        "F": 60,

        "G": 70,

        "H": 25
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    workbook.save(
        filepath
    )

    print()
    print(
        f"Excel report saved:"
    )

    print(
        filepath
    )


# ============================================================
# PROCESS ONE CONTACT
# ============================================================

def process_contact(
    page,
    contact
):

    name = contact["name"]

    phone = contact["phone"]

    message = contact["message"]

    result = {

        "name":
            name,

        "phone":
            phone,

        "message":
            message,

        "status":
            "FAILED",

        "error":
            "",

        "last_3_messages":
            [],

        "screenshot":
            "",

        "processed_at":
            ""
    }

    print()
    print("=" * 60)

    print(
        f"Processing: {name}"
    )

    print(
        f"Phone: {phone}"
    )

    print("=" * 60)

    try:

        # Open WhatsApp chat
        open_contact(
            page,
            phone
        )

        # Make sure chat/message box exists
        message_box = get_message_box(
            page
        )

        if message_box is None:

            raise RuntimeError(
                "Could not open the WhatsApp chat "
                "or message box was not found."
            )

        # Send personalized message
        send_message(
            page,
            message
        )

        # Screenshot
        screenshot = take_screenshot(
            page,
            name
        )

        result[
            "screenshot"
        ] = screenshot

        # Extract messages
        messages = (
            extract_last_three_messages(
                page
            )
        )

        result[
            "last_3_messages"
        ] = messages

        # Success
        result[
            "status"
        ] = "SENT"

        result[
            "processed_at"
        ] = datetime.now().isoformat()

        print()
        print(
            f"SUCCESS: {name}"
        )

    except Exception as error:

        result[
            "status"
        ] = "FAILED"

        result[
            "error"
        ] = str(error)

        result[
            "processed_at"
        ] = datetime.now().isoformat()

        print()
        print(
            f"FAILED: {name}"
        )

        print(
            f"Error: {error}"
        )

    return result


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print("PLAYWRIGHT WHATSAPP AUTOMATION")
    print("=" * 60)

    print()
    print(
        f"Project directory:"
    )

    print(
        BASE_DIR
    )

    print()

    # --------------------------------------------------------
    # Check Chrome
    # --------------------------------------------------------

    if not os.path.exists(
        CHROME_PATH
    ):

        raise FileNotFoundError(
            "Google Chrome was not found at:\n"
            f"{CHROME_PATH}"
        )

    print(
        "Google Chrome found:"
    )

    print(
        CHROME_PATH
    )

    # --------------------------------------------------------
    # Create Excel
    # --------------------------------------------------------

    create_contacts_file()

    # --------------------------------------------------------
    # Load contacts
    # --------------------------------------------------------

    contacts = load_contacts()

    if not contacts:

        print()
        print(
            "No contacts found in contacts.xlsx."
        )

        print(
            "Please add at least one contact."
        )

        return

    print()
    print(
        f"Loaded contacts: {len(contacts)}"
    )

    # --------------------------------------------------------
    # Start Playwright
    # --------------------------------------------------------

    with sync_playwright() as playwright:

        print()
        print(
            "Starting Google Chrome through Playwright..."
        )

        context = (
            playwright.chromium
            .launch_persistent_context(

                user_data_dir=
                    SESSION_DIR,

                headless=
                    HEADLESS,

                executable_path=
                    CHROME_PATH,

                viewport={
                    "width": 1366,
                    "height": 768
                },

                args=[
                    "--start-maximized"
                ]
            )
        )

        # ----------------------------------------------------
        # Get browser page
        # ----------------------------------------------------

        if context.pages:

            page = context.pages[0]

        else:

            page = context.new_page()

        # ----------------------------------------------------
        # Open WhatsApp
        # ----------------------------------------------------

        print()
        print(
            "Opening WhatsApp Web..."
        )

        page.goto(
            "https://web.whatsapp.com",
            wait_until="domcontentloaded"
        )

        # ----------------------------------------------------
        # Login
        # ----------------------------------------------------

        logged_in = (
            wait_for_whatsapp_login(
                page
            )
        )

        if not logged_in:

            print()
            print(
                "Login was not detected."
            )

            print(
                "Browser will remain open for 30 seconds."
            )

            page.wait_for_timeout(
                30000
            )

            context.close()

            return

        # ----------------------------------------------------
        # Process contacts
        # ----------------------------------------------------

        results = []

        for contact in contacts:

            result = process_contact(
                page,
                contact
            )

            results.append(
                result
            )

            # Random delay between contacts
            print()

            print(
                "Preparing for next contact..."
            )

            human_delay(
                2,
                5
            )

        # ----------------------------------------------------
        # Save reports
        # ----------------------------------------------------

        print()
        print("=" * 60)
        print("AUTOMATION FINISHED")
        print("=" * 60)

        save_json_report(
            results
        )

        save_excel_report(
            results
        )

        # ----------------------------------------------------
        # Summary
        # ----------------------------------------------------

        successful = sum(

            1

            for result in results

            if result["status"] == "SENT"
        )

        failed = sum(

            1

            for result in results

            if result["status"] == "FAILED"
        )

        print()
        print("=" * 60)
        print("SUMMARY")
        print("=" * 60)

        print(
            f"Total contacts : {len(results)}"
        )

        print(
            f"Sent           : {successful}"
        )

        print(
            f"Failed         : {failed}"
        )

        print()
        print(
            f"Screenshots:"
        )

        print(
            SCREENSHOT_DIR
        )

        print()
        print(
            f"Reports:"
        )

        print(
            REPORT_DIR
        )

        print()
        print(
            "Automation completed."
        )

        # ----------------------------------------------------
        # Close browser
        # ----------------------------------------------------

        context.close()


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    try:

        main()

    except KeyboardInterrupt:

        print()
        print(
            "Automation stopped by user."
        )

    except Exception as error:

        print()
        print("=" * 60)
        print("FATAL ERROR")
        print("=" * 60)

        print(
            error
        )

        print()